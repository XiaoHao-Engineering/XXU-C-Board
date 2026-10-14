# -*- coding: utf-8 -*-
"""生成《C 板工程 与 2027 哨兵参考工程 —— 底层与算法对比分析报告》DOCX。

用法：python gen_compare_report.py <输出路径.docx>

说明：所有对比结论均来自对两个工程源码的逐文件阅读，正文中带文件名/函数名/行号作为证据。
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = '微软雅黑'
MONO = 'Consolas'

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GOOD = RGBColor(0x1E, 0x7A, 0x34)
WARN = RGBColor(0xA6, 0x5B, 0x00)
BAD = RGBColor(0xB0, 0x2A, 0x2A)
GREY = RGBColor(0x55, 0x55, 0x55)


# --------------------------------------------------------------------------- 工具
def _set_rfonts(el, name):
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), name)


def set_style_font(style, name=FONT, size=None, bold=None, color=None):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    _set_rfonts(style.element, name)


def run_font(run, name=FONT):
    run.font.name = name
    _set_rfonts(run._element, name)


def para(doc, text='', style=None, bold=False, size=None, color=None, align=None,
         space_after=None, mono=False, italic=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        run_font(r, MONO if mono else FONT)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(doc, parts, style=None, align=None, space_after=None):
    p = doc.add_paragraph(style=style)
    for text, opt in parts:
        r = p.add_run(text)
        run_font(r, MONO if opt.get('mono') else FONT)
        r.bold = opt.get('bold', False)
        if opt.get('size'):
            r.font.size = Pt(opt['size'])
        if opt.get('color'):
            r.font.color.rgb = opt['color']
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        run_font(r)
        r.bold = True
    r = p.add_run(text)
    run_font(r)
    p.paragraph_format.space_after = Pt(2)
    return p


def set_col_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcpr.append(shd)


def cell_text(cell, text, bold=False, size=9, color=None, mono=False):
    """单元格文本；支持用 ** 包裹的加粗片段。"""
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    for i, seg in enumerate(str(text).split('**')):
        if seg == '':
            continue
        r = p.add_run(seg)
        run_font(r, MONO if mono else FONT)
        r.bold = bold or (i % 2 == 1)
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def no_split_rows(table):
    """禁止表格行跨页断开。"""
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))


def two_col_table(doc, rows, w1=3.3, w2=13.5, size=9):
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        cells = t.add_row().cells
        cell_text(cells[0], k, bold=True, size=size)
        shade(cells[0], 'F2F5F9')
        cell_text(cells[1], v, size=size)
    set_col_widths(t, [w1, w2])
    no_split_rows(t)
    return t


def item(doc, idx, title, ref, mine, gap, when, port):
    doc.add_heading('%d. %s' % (idx, title), level=3)
    two_col_table(doc, [
        ('参考工程怎么做', ref),
        ('本工作区现状', mine),
        ('具体差距', gap),
        ('什么工况下真的有优势', when),
        ('能否搬到 F407', port),
    ])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def big_table(doc, header, rows, widths, size=8.5, head_fill='1F4E79'):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, header):
        cell_text(c, h, bold=True, size=size, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, head_fill)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            cell_text(c, v, size=size)
    set_col_widths(t, widths)
    no_split_rows(t)
    return t


# --------------------------------------------------------------------------- 正文
def build(path):
    doc = Document()

    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)

    normal = doc.styles['Normal']
    set_style_font(normal, FONT, 10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    set_style_font(doc.styles['Title'], FONT, 21, bold=True, color=ACCENT)
    for lvl, sz in ((1, 15), (2, 12.5), (3, 11), (4, 10)):
        set_style_font(doc.styles['Heading %d' % lvl], FONT, sz, bold=True, color=ACCENT)

    # ================= 封面
    doc.add_heading('C 板工程 与 2027 哨兵参考工程', level=0)
    para(doc, '底层驱动 & 算法实现 对比分析报告', size=12.5, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=ACCENT)
    para(doc, '对比对象：① 本工作区 C-Board（RoboMaster 开发板 C 型 / STM32F407IGH6 / C11 / 纯 C）'
              '　② 参考工程 Robomaster_XinxiangUniversity_2027_Sentry'
              '（达妙 MC02 / STM32H723VGT6 / C++17）',
         size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
    para(doc, '结论依据：对两个工程源码逐文件阅读，正文均标注文件名 / 函数名 / 行号作为证据。',
         size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
    doc.add_paragraph()

    # ================= 一、结论摘要
    doc.add_heading('一、结论摘要（先看这里）', level=1)

    rich(doc, [
        ('一句话：', {'bold': True}),
        ('参考工程是「把单个控制器做深、把驱动做成可复用公共库、把踩过的坑写进注释」；'
         '本工作区是「把控制器做广（可任意组合）、把 3D 数学做全、把整车业务铺满」。'
         '两者不是同一块芯片上的同类东西，所以差异要先分成两类：', {}),
    ])

    big_table(doc,
              ['差异类别', '包含内容', '能不能借鉴'],
              [
                  ['① 算法与驱动设计',
                   '算法实现细节、设计模式、分层与命名规范、构建组织方式',
                   '★ 与 MCU 无关，可以直接借鉴 —— 本报告第 3 章逐条拆解'],
                  ['② MCU 平台特性',
                   'FDCAN / BDMA / MDMA / OCTOSPI / USB_OTG_HS / Cache / MPU / 多域 SRAM / 550MHz',
                   '✗ H7 专属，F407 上根本不存在，不能搬 —— 第 3 章全部单独标注'],
              ],
              [3.2, 9.8, 4.2])

    doc.add_paragraph()
    rich(doc, [('最值得马上抄的五件事：', {'bold': True})])
    bullet(doc, '把 PID 的工程化特性移植进 user_pid.c：显式控制周期 dt、微分先行、D 项低通、'
                '三段独立限幅、变速积分、积分分离、死区、前馈方向枚举、零位积分泄放、Clear()。'
                '全部与 MCU 无关，是唯一直接影响控制效果的部分。', '① ')
    bullet(doc, '给 UART/CAN 加"接收时间戳"，并补一个统一的 64 位微秒时间基准；'
                '现在的在线检测只有帧计数，判断不了"多久没来数据"。', '② ')
    bullet(doc, 'CAN 发送从主循环挪到定时器节拍里，发送时序由"尽力而为"变成"确定周期"。', '③ ')
    bullet(doc, '新增无效浮点数检测（NaN / Inf / 非正规数）并接到 IMU、DBUS、串口调参入口。'
                '这是典型的"平时没用、出事救命"的防御性代码。', '④ ')
    bullet(doc, '把 CMSIS-DSP 改成聚合编译单元（7 个 .c 代替 564 个），首次全量构建时间数量级下降。', '⑤ ')

    doc.add_paragraph()
    rich(doc, [('本工作区不该被比下去的地方：', {'bold': True}),
               ('控制器可组合的虚表接口与 7 种控制器、完整的四元数/旋转矩阵 3D 数学库、'
                '串口发送队列、功率限制、以及上层整车业务（三种底盘模式、云台锁世界、航向补偿）的完整度，'
                '都强于参考工程当前这份代码。详见第 四 章。', {})])

    # ================= 二、基本情况
    doc.add_heading('二、两个工程的基本情况', level=1)
    big_table(doc,
              ['项目', '本工作区 C-Board', '参考工程 2027_Sentry'],
              [
                  ['MCU / 内核', 'STM32F407IGH6（Cortex-M4）', 'STM32H723VGT6（Cortex-M7，单核）'],
                  ['主频 / 时钟', '168 MHz（HSE 12M）',
                   '550 MHz（HSE 24M，PLLM=3 / PLLN=68 / PLLFRACN=6144 / PLLP=1）'],
                  ['开发板', 'RoboMaster 官方开发板 C 型', '达妙 MC02 开发板'],
                  ['语言', 'C11，纯 C，全静态内存，无 new/delete', 'C++17（类 / 模板 / 引用 / 默认参数 / 运算符重载）'],
                  ['浮点', '软浮点（-mfloat-abi 默认 soft），CMSIS-DSP 从源码编',
                   '硬浮点（-mfpu=fpv5-d16 -mfloat-abi=hard）'],
                  ['分层', 'Core / User_Drives / User_Architect / User_Algorithm / User_Application',
                   'Core / User_File{1_Middleware, 2_Device, 5_Task} / USB_DEVICE'],
                  ['命名规范', 'user_xxx.c，按目录分层，文件名前缀不统一',
                   'alg_ / sys_ / drv_ / bsp_ / dvc_ / tsk_ 前缀，前缀强制单向依赖'],
                  ['注释', '中文行注释 + 关键处块注释（讲"怎么做"）',
                   '全工程 Doxygen 骨架 + 每函数每变量注释（@note 讲"为什么"与已知限制）'],
                  ['构建', 'CMake + Ninja，file(GLOB_RECURSE) 一把梭，564 个 DSP 编译单元',
                   'CMake + Ninja，显式源文件列表，INTERFACE/OBJECT/STATIC 三种库目标，'
                   'CMSIS-DSP 聚合为 7 个编译单元'],
                  ['调度', 'SysTick 任务表（任意周期 / 单次或重复 / 带参数），主循环跑 UART 队列',
                   '唯一 1ms TIM7 中断 + mod10/mod15/mod100 软件分片，主循环只放蜂鸣器'],
                  ['时间基准', 'HAL_GetTick()（1kHz）',
                   'TIM5 32 位 1MHz 计数 + 溢出计数 → 64 位微秒时间戳，并为每帧数据打时间戳'],
                  ['外设重点', 'CAN×2 / DBUS / HWT906 / BMI088(SPI) / 串口绘图 / JScope',
                   'FDCAN / USB CDC / OCTOSPI+W25Q64 / SPI6+WS2812 / 按键 / 蜂鸣器'],
                  ['业务完整度', '整车：三种底盘模式 + 云台锁世界 + 航向补偿 + 在线调参',
                   '当前这份代码是"公共库 + 单电机验证"：1 个 C620 + 波形发生器 + Serialplot'],
                  ['工程定位（自述）', '哨兵底盘控制框架',
                   'README 写明"公共库按长期复用、稳定优先、兼容优先维护"（共 8 行）'],
              ],
              [2.3, 7.0, 7.9])

    para(doc, '', space_after=2)
    rich(doc, [('⚠️ 重要前提：', {'bold': True, 'color': BAD}),
               ('参考工程跑在 H7 上，FDCAN、BDMA/MDMA、OCTOSPI、USB_OTG_HS、Cache/MPU、'
                '多域 SRAM 这些是芯片给的，不是代码写得好。本报告把它们单独标注为「H7 专属」，'
                '避免误当成可以搬过来的架构优点。另外需要澄清：', {})])
    bullet(doc, 'H723 是单核 Cortex-M7，工程里没有任何双核特性（无 HSEM / IPCC / OpenAMP 调用）——'
                'stm32h7xx_hal_conf.h 里的 HAL_HSEM_MODULE_ENABLED 只是 CubeMX 模板残留，'
                '不要把它当成架构优点。')

    # ================= 三、参考工程的优点
    doc.add_heading('三、参考工程的底层与算法优点（逐条拆解）', level=1)
    para(doc, '每条按「参考工程怎么做 / 本工作区现状 / 具体差距 / 什么工况下真的有优势 / 能否搬到 F407」'
              '五栏展开。', size=9.5, color=GREY)

    # -------- 3.1 PID
    doc.add_heading('3.1 控制器（算法核心）：单点做深', level=2)
    rich(doc, [('参考工程的 ', {}), ('Class_PID', {'mono': True}),
               ('（alg_pid.h / alg_pid.cpp）把现场踩过的坑几乎全做成了开关，'
                'Init() 共 16 个参数；本工作区的 ', {}), ('PID_Controller', {'mono': True}),
               (' 是 5 个参数的经典实现。下面按特性逐条拆。', {})], space_after=6)

    item(doc, 1, 'PID 有显式控制周期 dt（D_T）',
         'Init() 收 __D_T（默认 0.001s）并对 dt<=0 做除零保护（Init 时一次 + 每周期再查一次）；'
         'I 项按 Integral_Error += speed_ratio * D_T * Error 累加，D 项、F 项都除以 D_T。'
         '于是 K_I 的量纲是 输出/(误差·s)，K_D 是 输出/(误差/s)。',
         'Iout += ki * err；Dout = kd * (err - err[1])。没有 dt 参数，'
         '采样周期被隐含"吸收"进了 ki / kd。',
         '参考工程改控制周期只需改一个参数；本工作区一旦把 CHASSIS_PERIOD_MS 从 1 改成 2，'
         '所有已整定好的 ki / kd 全部失配，必须重新整定。',
         '任何需要改控制频率的场合：把 1kHz 降到 500Hz 省 CPU、为更高带宽提到 2kHz；'
         '以及串级控制内外环频率不同时。',
         '🟢 完全可搬（纯算法）。给 PID_Controller 加 dt 字段与 Init 参数，'
         'I 项乘 dt、D 项除 dt，存量参数按原周期换算一次即可。')

    item(doc, 2, '微分先行（D 作用在被控量的测量值上）',
         'Enum_PID_D_First 开关。开启后 D 项用 -(Now - Pre_Now)/D_T（并按 PID_REVERSE 处理符号），'
         '不再对误差微分，避免目标阶跃时 D 项冲击。',
         'D 项固定对误差微分：kd * (err - err[1])。',
         '目标值阶跃时误差同时跳变，D 项会产生一个巨大的尖峰（derivative kick）。'
         '按参考工程注释里的算例：Kd=0.5、阶跃 5rad、dt=1ms → 单拍尖峰 2500。',
         '目标值本身会阶跃的场合：云台锁世界时切换目标角、底盘模式切换、波形发生器给阶跃给定、'
         '遥控器摇杆快速到位。这时普通 PID 会"顶一下"甚至触发过流。',
         '🟢 完全可搬（纯算法）。改动量极小（一个 switch + 两个历史量）。')

    item(doc, 3, 'D 项一阶低通滤波（D_Filter_Alpha）',
         'Filtered_D_Out = alpha*raw + (1-alpha)*上一次，alpha 在 Init 和 Set_D_Filter_Alpha 两处'
         '都被 Constrain_Float 钳到 [0,1]。',
         '无任何滤波，D 项直接输出。',
         '编码器 / 陀螺的量化噪声会被 kd 放大后直接送进电流环，表现为电机"沙沙"响、发热。',
         '速度环靠编码器差分求速度（本工作区底盘 4 轮 M3508 正是如此）、或 D 项系数较大的场合。',
         '🟢 完全可搬（纯算法）。注意参考工程的 alpha 是采样域系数而非时间常数，'
         '同一套参数换采样率后等效截止频率会变；本工作区也可直接用已有的 user_fir.c 做更彻底的低通。')

    item(doc, 4, '三段独立限幅：积分项 / 微分项 / 总输出',
         'I_Out_Max、D_Out_Max、Out_Max 三个独立参数（0 = 不限制），D 在求和之前就被夹住，'
         '不会被后级 Out_Max 掩盖。另有一处关键细节：对积分累加器本身也按 I_Out_Max / K_I 钳位。',
         '只有 max_iout（积分限幅）与 max_out（总输出限幅），没有 D 项单独限幅。'
         '（积分累加器钳位这一条本工作区通过"直接钳 Iout"达到了等价效果。）',
         'D 项尖峰只能靠总限幅兜住，而总限幅通常按稳态需求设得较宽，所以尖峰仍会完整传到电流环。',
         '任何 D 项可能瞬间变大的场合（同第 2 条）。',
         '🟢 完全可搬（纯算法）。建议只补 D_Out_Max。')

    item(doc, 5, '变速积分（I_Variable_Speed_A / B 线性衰减）',
         '|误差| <= A 正常积分；A < |误差| < A+B 时积分速度线性衰减到 0；>= A+B 停止积分。'
         'B 是过渡带宽度，因为 ratio 在两处连续，所以不会产生跳变。',
         '积分速度恒定，只靠积分限幅兜底。',
         '大误差时积分照常累积，等误差变小后积分已经很大 → 超调与来回振荡。',
         '目标值大幅跳变后的恢复段：底盘模式切换、被外力推开后回位、初次上电追赶目标角。'
         '与"积分分离"相比它是连续的，输出不会在阈值处产生折点。',
         '🟢 完全可搬（纯算法）。')

    item(doc, 6, '积分分离（I_Separate_Threshold）',
         '误差超过阈值时把 Integral_Error 与 I_Out 清零（硬清零，不是冻结）。',
         '无。',
         '大误差段的积分全部白攒，只能靠限幅慢慢退掉。',
         '定点 / 定角控制的"快速靠近"阶段，例如云台大角度转向。',
         '🟡 思路可搬，但要改：参考工程是硬清零，'
         '对"靠积分顶重力"的轴（云台 pitch）在大阶跃瞬间会掉头/下沉 —— '
         '它自己的源码注释也承认 pitch 轴必须保留积分，前后矛盾。建议改成冻结或衰减。')

    item(doc, 7, '死区（Dead_Zone）',
         '|误差| < Dead_Zone 时把误差直接置 0，而且是在算 P **之前**清零，'
         '所以 P 与（未开微分先行的）D 一并被压掉。',
         'PID 本身无死区（本工作区只在应用层做遥控摇杆死区，位置不同）。',
         '执行机构在零位附近来回抖动（编码器量化 + 静摩擦造成的极限环），表现为电流纹波、嗡鸣、发热。',
         '云台 yaw 静止锁位、底盘静止保持。',
         '🟢 完全可搬（纯算法）。注意会牺牲一点静态精度，云台建议取很小值或不用。')

    item(doc, 8, '目标速度前馈（K_F）',
         'F_Out = K_F * (Target - Pre_Target)/D_T，直接加在总输出上（Out = P + I + D + F）。'
         '本质是"目标速度前馈"，不含重力常数前馈、不含加速度前馈。',
         '无。PID_Calculate() 虽然收了 sub_feedback 形参，但函数体里没有使用它。',
         '纯反馈天然滞后一拍；对斜坡 / 正弦给定，跟随误差与速度成正比。',
         '需要跟随动态目标的场合：云台跟随视觉目标、底盘跟随给定速度曲线，'
         '以及给波形发生器做前馈标定。',
         '🟢 完全可搬（纯算法）。本工作区的 user_parallel.c 也能从外部并联实现前馈，但内置更省事。')

    item(doc, 9, '方向枚举 PID_DIRECT / PID_REVERSE',
         '一个枚举把"控制量与被控量正相关还是负相关"表达清楚，并且误差取反、'
         '前馈取反、微分先行符号翻转三处保持一致。',
         '没有方向概念，需要调用者自己把目标或反馈取负。',
         '符号处理分散在调用点，容易漏、容易在改动时出错（典型 bug：云台 pitch 方向反了）。',
         '所有被控对象方向与控制量相反的场景：云台 pitch 轴、反向安装的电机。',
         '🟢 完全可搬（纯算法）。属于"低成本提升可读性与安全边际"的改动。')

    item(doc, 10, '零位积分泄放（ZPIB）',
         '目标与测量都落在死区内且开关打开时，不再累积新积分，并把已有 Integral_Error '
         '每周期乘 0.95 缓慢释放。源码注释还写明动机（静摩擦残留积分耗功发热）'
         '与反例（pitch 轴需积分顶重力、不能盲开）。',
         '无。',
         '电机静止时残留积分持续输出力矩、白白发热耗电。',
         '底盘轮电机静止保持（全向轮零位抖动）。注意 pitch 这类靠力矩抵消重力的轴不能开。',
         '🟢 完全可搬（纯算法）。建议只为底盘轮速环开启。'
         '另注意参考工程把触发条件写成了 |Target| < Dead_Zone，'
         '所以 Dead_Zone=0 时 ZPIB 永不触发 —— 两者是隐式耦合的。')

    item(doc, 11, '参数的在线可改 + 分项可观测 + 可复位',
         '15 个 inline Set_*（全部增益与限幅）；Get_P_Out / Get_I_Out / Get_D_Out / Get_F_Out / '
         'Get_Out / Get_Error / Get_Integral_Error 分项输出（实际被 Serialplot 采集）；'
         'Clear() 一键复位全部历史状态，电机掉线时会被调用。',
         '只有 Set_MaxOut 一个 setter；Pout / Iout / Dout 作为结构体字段可直接读'
         '（这点等价于 Getter，甚至更方便）；没有 Clear()，复位靠 memset 或手工改字段。',
         '在线调参只能改限幅；模式切换时无法优雅地"清掉历史状态"，切换瞬间容易踢一脚。',
         '上位机在线整定 PID（本工作区的 serialplot 调参表正需要全套 setter）；'
         '云台 / 底盘模式切换时清零积分。',
         '🟢 完全可搬（纯算法）。建议补齐全套 setter 与 PID_Clear()。')

    # -------- 3.2 基础数学库
    doc.add_heading('3.2 基础数学库：防御性工具与通用工具', level=2)
    rich(doc, [('需要先纠正一个容易误解的点：参考工程的 ', {}), ('alg_basic', {'mono': True}),
               (' 里', {}), ('没有', {'bold': True}),
               ('向量、矩阵、CRC、插值、独立一阶滤波、斜坡限制器 —— 它是一套"低层数值小工具"，'
                '与本工作区的 ', {}), ('user_coord.c', {'mono': True}),
               ('（3D 坐标 / 四元数库）大部分不重叠、互为补充。它有两类东西本工作区确实没有。', {})],
         space_after=6)

    item(doc, 12, '无效浮点数检测 Basic_Math_Is_Invalid_Float',
         '位级判定：把浮点 memcpy 到 uint32 后看指数域 —— ±0 合法；'
         '指数为 0 但尾数非 0 的次正规数**判为无效**；指数全 1 即 Inf / NaN 也判无效。'
         '不依赖 isnan / 浮点异常标志，中断里可放心调用。',
         '没有等价工具。user_coord.h 只定义了 PRECISION = 1e-6f 作为阈值常量，不做数值合法性判断。',
         '传感器通信出错（例如 BMI088 读回全 0xFF）、除零、姿态解算出现 NaN 时，'
         '脏值会一路流过 PID 直到电机电流指令，表现为电机突然满转或锁死；'
         '而且 NaN 一旦进入积分累加器就永久污染（NaN*0.95 仍是 NaN，限幅对 NaN 的比较全 false 也救不回）。',
         '任何"安全相关"的数据链路上。属于典型的"平时没用、出事救命"的防御性代码。',
         '🟢 完全可搬（纯逻辑，几行）。建议在 IMU 采样、DBUS 解析、串口调参入口各加一道。')

    item(doc, 13, '通用数值工具集',
         'Basic_Math_Modulus_Normalization(x, modulus) —— 通用取模归一化（不限于角度，'
         '内部用 fmodf(x+mod/2, mod) 并补偿负值，正确避开了 fmodf 返回负数的坑）；'
         'Basic_Math_Endian_Reverse_16 / 32 —— 大小端翻转，且提供"原地交换"与'
         '"返回值并可选写出"两种重载；Basic_Math_Sum_8/16/32 —— 数组累加校验和；'
         'Basic_Math_Float_To_Int / Int_To_Float —— 浮点/整型区间线性映射（含退化区间保护）；'
         'Basic_Math_Sinc；以及 BASIC_MATH_RPM_TO_RADPS 等 const float 具名单位换算常量'
         '（不是宏，有类型检查）。',
         'user_coord.h 有 Math_WrapAngleRad / Math_WrapAngleDeg（只针对角度）、'
         'Math_Deg2Rad / Math_Rad2Deg、FastInvSqrt；端序与校验和则散落在 '
         'user_serialplot.c、user_dji_bus.c 内部各自实现。',
         '角度归一化只覆盖角度（编码器计数等其他模数就要另写）；'
         '校验和 / 端序没有公共实现，重复代码分散在各协议文件里。',
         '新增协议解析（裁判系统、视觉通信、外部 Flash 参数帧）、编码器多圈处理。',
         '🟢 完全可搬（纯逻辑）。建议收敛到 user_coord.c 或新建 user_math_util.c。')

    item(doc, 14, 'C++ 模板化通用函数',
         'template<typename Type> Basic_Math_Constrain(Type x, Min, Max) 传值版，'
         '以及指针重载版（就地修改 + nullptr 检查）；Basic_Math_Abs(Type)。'
         '一套代码覆盖 float / int / uint 所有类型，且注释里诚实写明'
         '"对无符号类型无意义、对有符号最小负数有溢出风险"。',
         '每处需要限幅就手写 if-else（user_pid.c、user_controller 家族、user_serialplot.c 里都有）。',
         '重复代码多、类型漏改有风险；但纯 C 没有模板。',
         '属于工程整洁度问题，不直接影响控制效果。',
         '🔴 语言特性，纯 C 用不了。可用 _Generic 宏或按类型各写一份替代。性价比低，不必强求。')

    # -------- 3.3 波形发生器
    doc.add_heading('3.3 波形发生器：从"测试件"升级为"系统辨识激励台"', level=2)

    item(doc, 15, 'Class_Waveform：13 种参数化波形',
         '支持 HOLD / STEP / SQUARE / SINE / TRIANGLE / SAW / RAMP / PULSE / DOUBLET / '
         'CHIRP（线性扫频）/ PRBS / NOISE / RELAY 共 13 种。'
         '实现上有几处扎实的细节：换频时相位连续（相位累积 Wrap_01(Phase + Freq*DT)）；'
         'CHIRP 相位连续线性扫频；PRBS 用 LFSR 且位定时用 while 追赶（调用慢于位率也不丢位）；'
         'RELAY 带滞环。参数化 + 出口统一限幅 + 时间推进集中在一处。',
         '有 user_waveform.c（WAVEFORM_DRIVES），在 main.c 里作为"PID 测试用波形发生器"创建，'
         '定位是调试辅助，没有形成"给已知激励 → 采集响应 → 拟合参数"的闭环用法。',
         '缺少一套标准工作流：本工作区已经有 JScope 和 Serialplot 可以采集响应，'
         '缺的是把波形激励与辨识流程串起来。有了它，PID / ADRC 的参数可以先算再调，'
         '而不是纯手拧。',
         '整定 ADRC / LADRC 的 b0 与观测器带宽、标定前馈系数、验证速度环带宽、'
         '用 DOUBLET 辨识摩擦与齿隙。',
         '🟢 机械结构可搬（波形算法与 MCU 无关）；本工作区已有等价基建'
         '（波形 + JScope + Serialplot），属于"补工作流"而非"补代码"。')

    item(doc, 16, '⚠️ 一个必须澄清的落差：自整定并不存在',
         '波形库里已经备好了 RELAY（继电反馈，Åström–Hägglund 自整定所需的激励原语）'
         '以及 PRBS / CHIRP 等辨识激励，但目前这份接入代码用的是 0.5Hz / ±5rad 的正弦角激励。',
         '本工作区同样没有自整定。',
         '要点是：参考工程全工程检索不到 Relay() 的任何调用点，也没有计算极限环周期 Tu / '
         '幅值 Ku 或 Z-N / Tyreus-Luyben 参数的代码。',
         '如果需要自动整定，两边都得从头做。',
         '🟡 激励原语可搬，整定算法两边都没有 —— '
         '不要因为看到 RELAY 就以为"参考工程支持自整定"；写文档和对外介绍时都要避免这个误解。')

    # -------- 3.4 驱动层
    doc.add_heading('3.4 驱动层：统一风格与确定性', level=2)

    item(doc, 17, '回调注入 + 中断/DMA 全异步解耦，ISR 文件零业务代码',
         '每个外设一个为全局唯一实例（CAN1/2/3_Manage_Object、UART1..10_Manage_Object、'
         'SPI1..6_Manage_Object、TIM1..24_Manage_Object、OSPI1/2_Manage_Object、'
         'USB_Manage_Object），结构体内存驱动句柄 + 用户回调指针；Init() 时注入回调。'
         'stm32h7xx_it.c 里只有 HAL_xxx_IRQHandler(&handle)，全工程 15 处 HAL 弱符号'
         '（HAL_FDCAN_RxFifo0Callback、HAL_UARTEx_RxEventCallback、HAL_SPI_TxCpltCallback、'
         'HAL_TIM_PeriodElapsedCallback 等）被驱动层统一接管。'
         '→ 中断 → 驱动 → 用户回调，三层单向，业务代码完全不碰 HAL 回调。',
         '本工作区也是"驱动层接管 HAL 回调 + 用户注册回调"的结构'
         '（user_can.c 里覆写 HAL_CAN_RxFifo0MsgPendingCallback、'
         'user_uart.c 里覆写 HAL_UARTEx_RxEventCallback，再遍历回调数组），'
         'stm32f4xx_it.c 同样只调 HAL_xxx_IRQHandler。'
         '差别在实例管理方式：参考工程是**编译期固定 10 个 UART / 3 个 CAN 全局对象**，'
         '靠手写 if-else 链（UART_Get_Manage_Object）反查；'
         '本工作区是**运行时注册表**（uart_drives[UART_NUM]，UART_Init 时入表），'
         '但有 UART_NUM=5、CAN_NUM=2 的硬上限。',
         '两种方式各有取舍：参考工程静态、可预测、无上限但加一路外设要改 2~3 条 if-else 链；'
         '本工作区加实例只改一个宏，但数组是固定上限。本质差距不大。',
         '这一条是本工作区已经做对的地方，不需要改 —— 列出来是为了说明"分层解耦"不是参考工程的独有优点。',
         '🟢 已具备，无需移植。')

    item(doc, 18, 'init_finished 全局初始化闸门',
         'extern bool init_finished 在 CAN / UART / USB / TIM 四个驱动的回调入口被检查，'
         'Task_Init() 末尾才置真。作用是屏蔽"外设中断已使能、但软件层还没初始化好"这段窗口，'
         '避免初始化期间乱回调。语义针对性各不相同：CAN 是"只清 FIFO 不回调"，'
         'UART 是"只重启接收不回调"。',
         '本工作区的做法是**顺序控制**：CubeMX 的 MX_*_Init 先跑完，'
         '再在 USER CODE 2 里逐个调 CAN_Init / UART_Init，'
         '而 SysTick 任务表用 SysTick_StartTask() 显式启动，未启动的任务不会被调用。'
         '功能上等价（初始化期间不会有业务回调），但没有一个统一的"就绪"标志。',
         '本工作区的机制是隐式的，依赖"初始化顺序正确"这一约定；'
         '参考工程是显式的，任何人新增驱动都能看到这个闸门。',
         '外设初始化比较慢、或中途可能失败需要跳过后续业务时更有价值。',
         '🟢 可搬（纯逻辑），但属于风格改进而非功能缺口。'
         '注意参考工程这个标志是普通 bool 却在 ISR 里读 —— 严格来说是数据竞态，'
         '移植时应写成 volatile。')

    item(doc, 19, '每帧数据带接收时间戳',
         'CAN 的 Rx_Timestamp、UART 的 Rx_Time_Stamp、SPI / USB / OSPI 的 Timestamp，'
         '都在驱动层的接收完成点用 SYS_Timestamp.Get_Current_Timestamp() 打好，'
         '设备层不必自己取时间。',
         '本工作区没有接收时间戳。在线检测靠帧计数：DBUS 用 rx_count、'
         '电机用 Flag 计数在 100ms 任务里比对，变化即在线。',
         '帧计数只能回答"有没有新数据"，回答不了"多久没来了"。'
         '做超时判断、丢帧统计、帧间隔抖动分析都需要时间戳。',
         '调试通信不稳（DBUS 偶尔丢帧、CAN 总线干扰）、做"数据陈旧则降级"的安全策略、'
         '统计各总线的实际帧率。',
         '🟡 思路可搬，但建议连时间基准一起补：本工作区目前只有 HAL_GetTick()（1kHz、1ms 分辨率），'
         '要微秒级就得再占一个 32 位定时器（F407 的 TIM2 / TIM5 也是 32 位，做法一致）。')

    item(doc, 20, 'SPI 片选参数化 + DMA 回调里释放',
         'SPI_Transmit_Data(...) 把 CS 的 GPIOx / Pin / Level 作为**每次调用的入参**传入，'
         '驱动在执行前拉有效电平，在 TxCplt / TxRxCplt / Error 回调里释放'
         '（失败路径也补一次释放）。同一条 SPI 挂多个从设备时可以自由换片选。',
         '本工作区用 SPI_DRIVES 结构体把 cs_port / cs_pin 固化成实例状态'
         '（user_spi.c 的 SPI_Init 一次绑定），SPI_Transfer 内部管理片选。'
         'BMI088 用两个 SPI_DRIVES 实例（user_spi_accel / user_spi_gyro）指向同一个 hspi1、'
         '分别持有 PA4 / PB0，同样实现了"一条总线多从设备"。',
         '都能满足多从设备需求。参考工程更灵活（同一次调用可换不同从设备），'
         '本工作区更省心（绑定一次，调用处不用管片选）。',
         '同一条 SPI 挂多个不同片选时序要求的从设备时，参数化更有优势。',
         '🟢 已具备等价能力，无需移植。')

    item(doc, 21, '⚠️ 参考工程驱动层反而弱于本工作区的两点',
         '① **没有发送队列**：UART 是单槽 512B + HAL_BUSY、USB 是 USBD_BUSY、'
         'CAN 是 Tx FIFO 满则 HAL_BUSY；而且上层一律忽略返回值'
         '（dvc_serialplot.cpp、dvc_vofa.cpp、dvc_motor_dji.cpp 的发送调用都没看返回值）'
         '→ 高负载下静默丢帧、丢控制帧，无计数无告警。'
         '② 驱动层**没有超时、没有错误码、没有日志、没有断言**：'
         '统一复用 HAL_StatusTypeDef 且多数 Init 返回 void；'
         'USE_FULL_ASSERT 全工程未定义，assert_param 被编译掉；'
         'Error_Handler() 只有 __disable_irq(); while(1)。',
         '① 本工作区的 UART 发送是**队列化**的（user_uart.c 的 tx_queue，'
         'Queue_Push 拷贝进节点、失败时 Queue_FreeNode 归还，能连续排队多帧）。'
         '② 本工作区同样没有统一的错误码体系，但 user_uart.c 明确处理了'
         '"发送失败要把节点释放掉，否则队列会堵死"这类问题，注释写明了原因。',
         '在"连续发多帧"这个场景上，本工作区的发送队列明显更稳；'
         '在错误可观测性上，两边都不足。',
         '串口绘图在 1kHz 任务里刷帧、DBUS 转发、CAN 连续下发时。',
         '🟢 本工作区在这两点上更优，**不要**把参考工程的无队列设计"优化"过来。')

    # -------- 3.5 设备层（电机）
    doc.add_heading('3.5 设备层（DJI 电机）：几个真正精彩的设计', level=2)

    item(doc, 22, '按"电调型号"建模 + 偏移指针式共帧',
         '三个并列类（Class_Motor_DJI_GM6020 / Class_Motor_DJI_C610 / Class_Motor_DJI_C620），'
         '抽象轴是**电调**而不是电机（M2006 归 C610、M3508 归 C620，电流标定系数只由电调决定，'
         '这对 DJI 生态是正确的抽象）。核心是 allocate_tx_data()：在 Init 时**只算一次指针**，'
         '把每个电机指向 12 个静态 8 字节帧缓冲（CAN1/2/3 × {0x1FE, 0x1FF, 0x200, 0x2FE}）中'
         '属于自己那 2 字节的地址，并精确复刻 DJI 协议（0x1FE / 0x2FE 是 GM6020 专用；'
         '0x201~0x204 走 0x200、0x205~0x208 走 0x1FE/0x1FF、0x209~0x20B 走 0x2FE）。'
         '→ 电机对象只多占 2 字节指针，Output() 只写自己那 2 字节，'
         '**无需拼帧、无需互斥**；型号与 ID 段不匹配时 Tx_Data = nullptr，'
         'Output() 开头判空静默返回，不会误写别人的帧。',
         '本工作区是单一 DJI_MOTOR_DRIVES 结构体 + 类型枚举（M3508_gear / GM6020）+ '
         '控制模式枚举，由 user_dji_motor.c 统一组帧发送八字节报文。'
         '控制器是**注入**的（DJI_Motor_Init 传 CONTROLLER_INTERFACE* 指针）。',
         '参考工程的共帧设计更省（无拼帧开销、天然线程安全）；'
         '本工作区的控制器注入更灵活（见第 4 章）。'
         '另外本工作区的电机 ID 是**裸数字**：'
         'DJI_Motor_Init(&user_wheel_fl, &user_can_1, 2, ...) 里的 2、3、1、4 是魔数，'
         '而参考工程用 Enum_Motor_DJI_ID 类型化枚举。',
         '多电机高频率下发（本工作区 4 轮 + 云台共 5 个电机，1kHz 下发）时，'
         '零拼帧开销与线程安全是有意义的。',
         '🟡 共帧设计需改造（本工作区已有可用的组帧实现，收益主要是省 CPU 与可读性）；'
         '**ID 枚举化是低成本高收益**，建议直接做。')

    item(doc, 23, '首帧编码器基准保护（Encoder_Init_Flag）',
         '三个电机类都有这个标志。注释写明动机："第一帧只建立编码器基准，'
         '防止 Pre_Encoder 默认 0 导致误判跨圈"。'
         '没有这段代码，上电时若电机恰好停在编码器 6000 处，'
         '第一帧就会被判成反向跨半圈，Total_Round 立即 −1，角度直接跳 2π。',
         '本工作区的 dji_motor 同样需要处理多圈角度，'
         '（user_dji_motor.c 有转子角度与多圈累加逻辑）—— 但本次未逐行核对该保护是否等价。',
         '这是"上电瞬间角度跳一圈"这类偶发怪现象的根因之一，'
         '属于非常典型的、只有长期调试才会暴露的坑。',
         '任何使用增量式编码器做多圈角度累积的场合。',
         '🟢 必须确认本工作区有等价保护；若没有，属于**潜在 bug**，应优先补齐。')

    item(doc, 24, '掉线时主动清零共享帧里的残留电流',
         'TIM_100ms_Alive_PeriodElapsedCallback() 比对 Flag == Pre_Flag 判断在线；'
         '掉线时 TIM_Calculate_ 一进来就 Clear() 并 return。'
         'Clear() 做了 6 件事，其中最容易被忽略、也最关键的是：'
         '**if (Tx_Data) { Tx_Data[0] = Tx_Data[1] = 0; }** —— '
         '不仅自己不输出，还把共享帧里自己那 2 字节清零，'
         '避免"上一次的残留电流指令"继续被别的电机带出去。',
         '本工作区有离线检测与停机处理（电机离线时清零输出），'
         '但需要确认共享帧场景下的残留清理是否等价。',
         '共帧设计下，一个电机掉线后如果不主动清零自己那 2 字节，'
         '其他电机照常发送时会把它的旧电流值一起发出去 —— 这是共帧方案特有的陷阱。',
         '多电机共用一帧时任意一个掉线 / 初始化失败的场合。',
         '🟢 可搬；若本工作区的组帧是"每次重新填满 8 字节"，则天然不存在此问题，'
         '需按实际实现确认。')

    item(doc, 25, '外部反馈覆盖 + 单次有效语义',
         'Set_Now_External_Angle / Set_Now_External_Omega 除了写值还置一个 Flag；'
         'PID_Calculate() 开头做 float feedback = Flag ? Now_External_Angle : Rx_Data.Now_Angle，'
         '然后**用完即清 Flag**。'
         '上层可以每周期"喂"一次真实反馈（IMU、外部编码器），不喂就自动回退到电机自带编码器，'
         '**不需要改动类结构**。',
         '本工作区通过 CONTROLLER_INTERFACE 注入控制器，'
         '反馈值由调用方在 Calculate(controller, main_feedback, sub_feedback) 里传入，'
         '本身就支持任意反馈源 —— 机制不同但能力等价，甚至更直接。',
         '两者都能解决"减速箱回差 / 编码器装在电机侧导致精度不够"的问题。'
         '本工作区的做法反而更干净（反馈由调用方决定，不需要"单次有效"这种状态机语义）。',
         '需要融合 IMU 与编码器、或做零位闭环时。',
         '🟢 本工作区机制已更优，无需移植。')

    item(doc, 26, '⚠️ 参考工程电机的两个功能缺口（本工作区反而更强）',
         '① **功率限制完全没实现**：POWER_MODEL 四个系数全为 0，'
         'power_calculate() 按 K0*I*ω + K1*ω² + K2*I² + A 算出来恒为 0，'
         'Get_Now_Power() 永远返回 0。唯一限幅是 Constrain(&Out, -16384, 16384)'
         '（电调量程，不是功率）。'
         '② **没有任何温度判断**：Data_Process 里没有 if (Temperature > ...)，'
         '温度保护缺失。',
         '本工作区有功率限制：DJI_Motor_Set_Power_Limit()，'
         '并在 chassis_config.h 里定义了 CHASSIS_WHEEL_POWER_MAX（单轮功率上限）等参数，'
         'main.c 里对云台 GM6020 也调用了 Set_Power_Limit。',
         '这是**功能性缺口**而非优化项：哨兵超功率会被裁判系统惩罚，超温会烧电机。'
         '参考工程这一版把公式形态预留了但没标定参数。',
         '整车功率受限的场合（哨兵、英雄等有功率上限的兵种）。',
         '🟢 本工作区已具备，**不要**因为"参考工程是公共库"就以为它更完备 —— 这一点它更弱。')

    # -------- 3.6 上位机通信
    doc.add_heading('3.6 上位机通信与绘图', level=2)

    item(doc, 27, 'Set_Data 存地址不存值，一次注册永久生效',
         'Class_Serialplot 的 Set_Data(int Number, ...) 用 va_arg 把**变量地址**'
         '存进 const void *Data[24]，头文件明确注释"存的是变量地址、不是当前值；'
         '每次 TIM_1ms_Write 才读最新值"。Output() 时按类型 memcpy 现取。'
         '→ 用户只需在 Task_Init 里写一次 Set_Data(8, &a, &b, ...)，之后永远不用管，'
         '而且保留原始位模式精度（不像格式化字符串会损失精度）。',
         '本工作区同样是地址注册：Serialplot_Set_Data(&user_serialplot, 5, '
         '&user_gimbal.yaw, &user_gimbal.target_yaw, ...)。机制一致。',
         '没有差距，这一条两边都对。',
         '—',
         '🟢 已具备，无需移植。')

    item(doc, 28, '绘图协议可选数据类型 + 可开关校验',
         'Serialplot 支持 8 种数据类型（UINT8/16/32、INT8/16/32、FLOAT、DOUBLE），'
         '由 Enum_Serialplot_Data_Type 配置，Get_Data_Type_Size() 映射字节数；'
         '校验和可开关（默认开）。另有 Vofa 模块提供第二套协议'
         '（justfloat 帧尾 0x7F800000 + 累加和，需要上位机侧 RMFloatSum8.dll），'
         '但**当前一个 VofaSum8 对象都没创建，是"完备但零接入"的库**。',
         '本工作区是固定 float32 + 0xAB 帧头 + 1 字节累加和的自定义协议，'
         '通道数由 CHASSIS_SERIALPLOT_MODE 编译期选择。',
         '参考工程协议灵活性更高（可选整型省带宽、可关校验）。'
         '本工作区协议固定但够用，而且**在线调参表带 min/max 钳位**'
         '（SERIALPLOT_VAR 结构体），安全性上反而更好。',
         '需要传整型状态量（模式、错误码）省带宽时；'
         '或需要与第三方上位机（VOFA+）对接时。',
         '🟡 可选数据类型可搬（纯协议逻辑，改动小）；'
         'Vofa 协议可搬但价值取决于你是否用 VOFA+。')

    item(doc, 29, '⚠️ 参考工程 USB 绘图发送存在数据竞争（借鉴时的反面教材）',
         'CDC_Transmit_HS() 里只做 USBD_CDC_SetTxBuffer(&hUsbDeviceHS, Buf, Len) —— '
         '**只记指针不拷贝**；而 Serialplot / Vofa 的 USB 版每 1ms 就 Output() 覆写同一块 Tx_Buffer，'
         '且返回值被完全忽略（USBD_BUSY 被吞）。全速 USB 下 40 字节帧 + 主机轮询可能 > 1ms，'
         '会间歇发出撕裂帧（校验和能发现但没人看）。'
         '对比之下 UART 版有 gState 忙检查，USB 版没有 —— 两条路径保护不对称。',
         '本工作区的串口绘图走 UART（USART1）+ DMA，'
         'user_uart.c 的发送队列只在 gState == HAL_UART_STATE_READY 时取下一帧，'
         '有明确的忙/闲判断与失败回退。**没有这个问题。**',
         '本工作区的发送路径更安全。',
         '—',
         '🟢 本工作区更优，无需移植。')

    # -------- 3.7 BSP
    doc.add_heading('3.7 BSP 层：几个值得借鉴的外设设计', level=2)

    item(doc, 30, '按键：静态侵入式链表，加按键 = 加一行代码',
         'Class_ArkKey 用 static Class_ArkKey *head/tail + 成员 next 组成侵入式链表，'
         'Init(GPIOx, Pin) 时尾插。所有批量操作（ClearAllFlags / ScanKeyState / UserApp / '
         'TIM_Update_PeriodElapsedCallback）都是一次遍历。'
         '调用方只需在 15ms 时间片里写一句 Class_ArkKey::TIM_Update_PeriodElapsedCallback()，'
         '**新增按键只加一行 Init**，调度代码一字不改。'
         '消抖用"连续两次采样一致"的时间采样法（非计数法），注释诚实写明'
         '"只能滤除小于 15ms 的抖动；连续按下间隔需 > 30ms"。',
         '本工作区**没有按键支持**（C 板上也没有用户按键的既定用途）。',
         '差距是"有没有"而非"好不好"。',
         '需要板载交互（拨弹、模式切换、急停复位、免上位机改参数）时非常有用。',
         '🟢 可搬（纯 C 也能照这个思路做，只是没有类，用结构体 + 指针数组即可）。'
         '但 F407 C 板上是否有可用按键引脚需要先确认。')

    item(doc, 31, '蜂鸣器：两字段谱面结构 + 可读的音名常量',
         '曲库用 typedef struct { float frequency; float beat; } 两个字段表示一个音符'
         '（frequency = 0 表示休止符，beat 是拍数，允许小数做装饰音）；'
         '配合 SONG_NOTE(freq, beat) / SONG_REST(beat) 两个宏，'
         '谱面直接写成可读文本。117 个音名常量（BUZZER_FREQUENCY_A0 ... C8，'
         '含升号/降号别名）逐个定义为 const float，写谱不用查表。'
         'PWM 换音用"先关定时器 → 清 CNT → 改 ARR/CCR → 手动 GenerateEvent(UPDATE) '
         '立即装载影子寄存器 → 再开"的流程，保证换音不出现半周期怪音。',
         '本工作区有 user_startup_music.c + User_Lib/user_music.h（dji_starting_music），'
         '结构上要简单一些（开机音乐单曲）。',
         '参考工程在"多曲库 + 换音无怪音"上更成熟。',
         '需要多首提示音、报警音、状态音，以及换音时不允许出现杂音时。',
         '🟡 谱面结构与换音流程可搬；117 个常量可搬但占 Flash，'
         '更适合做成数组按索引取（参考工程没做数组化，扫频要自己写 switch）。'
         '⚠️ 注意参考工程的 Play_Note 用 HAL_Delay 阻塞，在主循环里播一首曲子会卡住主循环数十秒 —— '
         '**这一点不能照抄**，正确做法是做成 1ms tick 驱动的 {freq, duration} 状态机。')

    item(doc, 32, 'WS2812：SPI + DMA 位编码，连 Reset 也走 DMA（零阻塞）',
         '把每个 WS2812 bit 编码成一个 SPI 字节（0x60 / 0x78），'
         'SPI6 用 HSE 24MHz ÷ 4 = 6MHz → 1 bit = 166.7ns、8 bit = 1333ns，'
         '正好落在 WS2812B 的 1250±600ns 窗口内；'
         'Reset 也用 240 个 0 字节由 DMA 送出（≈320µs > 280µs 要求）。'
         '整个 Write() 零阻塞，与蜂鸣器的 HAL_Delay 形成鲜明对比。'
         '驱动内部做 RGB→GRB 重排，对外 API 仍是 RGB，层次干净。'
         '缓冲区按 DMA 域分开（SPI6 挂 BDMA，缓冲放 RAM_D3），'
         '这是 H7 上最容易做错的细节，它做对了。',
         '本工作区没有 WS2812（C 板是三色 GPIO 灯）。',
         '这是"参考工程多出来的外设能力"，不是本工作区的短板。',
         '需要状态灯带 / 灯效（血量、状态、调试可视化）时。',
         '🟡 思路可搬（F407 用 SPI+DMA 做 WS2812 同样可行），'
         '但属于新功能而非移植；⚠️ 注意该模块头文件注释里的 0xC0/0xF0 与代码实际的 0x60/0x78 '
         '不一致（代码是对的，注释是错的），照注释"修正"会引入半时钟相位错误。')

    item(doc, 33, 'W25Q64 外部 Flash：全异步状态机 + 硬件自动轮询 + 寄存器层独立成文件',
         '四层职责清晰：应用层 → 器件层（状态机）→ **寄存器层独立成文件**'
         '（bsp_w25q64jv_register.h 只有 8 条指令枚举，换型号只改这一个文件）→ 总线层。'
         '状态机由 OSPI 回调驱动，用 HAL_OSPI_AutoPolling_IT（硬件轮询）'
         '而不是软件 while(BUSY) —— 这是"不阻塞 CPU"的关键。'
         '初始化有 4 步握手：读 JEDEC ID 校验（0x00EF4017）→ 读状态寄存器 2 → '
         '若 QE 未置位则写使能 + 轮询 WEL + 写状态寄存器 2 打开四线模式 → '
         '**自动开 QE 很贴心**（新片子出厂 QE=0，不做这步四线读必然失败，很多工程现场踩坑）。'
         '边界检查扎实：长度、地址范围、**跨页检查**、4KB 对齐、忙互斥、未初始化检查。'
         '异步 API 里先 memcpy 保存用户数据，防止异步期间原缓冲被改 —— 异步设计的正确做法。',
         '本工作区没有外部 Flash（A 板也没有；这是参考工程多出来的能力）。',
         '同样是"参考工程多出来的外设能力"。'
         '⚠️ 但它**没有任何磨损均衡 / 坏块管理 / 掉电保护**，'
         '而且自动轮询**没有软件超时**：Flash 若一直 BUSY，Get_Busy_Flag() 会永久为真，'
         '后续所有操作永久 HAL_BUSY，应用层无法区分"暂时忙"与"永久坏"。',
         '需要掉电保存参数（PID 参数、零位标定值、场地信息）时。',
         '🟡 架构可搬（F407 用 QUADSPI + DMA 实现同样的异步状态机），'
         '但移植时必须补软件超时与磨损考虑。')

    # -------- 3.8 任务与时间
    doc.add_heading('3.8 任务与时间基准', level=2)

    item(doc, 34, 'main.c 极简：用户代码只有 Task_Init() / Task_Loop()',
         'main.c 里 CubeMX 生成部分之外，用户代码只有两行：USER CODE 2 里的 Task_Init()，'
         '和 while(1) 里的 Task_Loop()。**全部业务在 User_File/5_Task/tsk_config_and_callback.cpp 里**，'
         '而 User_File 目录与 CubeMX 生成物物理隔离'
         '（ProjectManager.KeepUserCode=true、CoupleFile=true）。'
         '→ 在 CubeMX 里重新 GENERATE CODE，绝对不可能伤到业务代码。'
         'Task_Init() 同时是一块可读性极高的"接线板"：'
         '十几行把每个驱动与它的回调一一对应，一眼看清全系统的连接关系。',
         '本工作区的 main.c 就是业务初始化本体（USER CODE 2 里 60 多行：'
         'JScope / TIMER / UART / LED / 蜂鸣器 / CAN / 4 个底盘电机 + PID / 底盘对象 / '
         'DBUS / HWT906 / BMI088 / 姿态 / 串口绘图 / 云台 / 底盘控制任务 / 灯任务）。'
         '这些都靠 USER CODE 块保护 —— 本次移植过程中已经实测过：'
         'CubeMX 重新生成确实保住了 USER CODE 块，但这也是"依赖生成器守规矩"。',
         '参考工程的方案对"重新生成"是**结构性免疫**，本工作区是**约定性免疫**。'
         '本次实测证明约定有效，但 main.c 每次生成都要仔细核对，心智负担更大。',
         '频繁改 CubeMX 配置、需要反复重新生成代码的项目。',
         '🟢 完全可搬，且性价比很高：把 USER CODE 2 的内容搬进一个 '
         'User_Application/app_init.c 的 App_Init() / App_Loop()，main.c 只留两行。')

    item(doc, 35, '唯一 1ms 硬周期 + mod 软件分片',
         '所有周期任务都在 TIM7 的 1ms 中断里，用 static 局部计数器做时间片：'
         '1ms 跑按键标志 / 波形更新 / Serialplot 组帧 / PID 计算 / CAN 组包发送；'
         '10ms 刷 WS2812；15ms 扫按键；100ms 做电机存活检测。'
         'modN == N 用比较代替取模（避免除法），且天然对齐 1ms 网格。'
         '优点：只有**一个**硬件定时器、任务表集中在一个函数里、极好读。'
         '主循环只有一句蜂鸣器播放，形同虚设。',
         '本工作区用 SysTick 任务表：SysTick_InitTask(&task, arg, delay, period, mode, callback) '
         '+ SysTick_StartTask()，支持任意周期、单次 / 重复、带参数、可停止。'
         '底盘控制任务 1kHz、LED 闪烁任务、BMI088 采样任务、HWT906 等都挂在表上；'
         '主循环跑 UART 收发队列与 CAN 发送队列（loop_event 表）。',
         '两者都是无 RTOS 的软件调度。本工作区的任务表**更通用**'
         '（任意周期、单次任务、带参数），参考工程的写法**更集中更易读**'
         '（所有周期任务一眼看全）。'
         '关键差别在**控制与发送的所在地**：参考工程 PID 与 CAN 发送都在 1ms 中断里'
         '（时序确定）；本工作区 CAN 发送在主循环 loop_event 里（尽力而为，会被主循环耗时影响）。',
         'CAN 下发时序要求严格（多电机同步、与云台控制对齐）时。',
         '🟡 把 CAN 发送从主循环挪进定时器节拍即可，改动很小；'
         '任务表本身不需要改（本工作区的更通用）。')

    item(doc, 36, 'TIM5 32 位 1MHz → 64 位微秒全局时间基准',
         '用一个 32 位定时器（TIM5）跑 1MHz（PSC=275-1，ARR=3600000000-1，'
         '更新中断 1 小时一次），中断里给 volatile uint32 溢出计数 ++；'
         '读取时用"读溢出 → 读 CNT → 再读溢出，不等则重读"的循环避免跨边界撕裂，'
         '最终返回 uint64 微秒。另外提供 Class_Time（int64 微秒）做时间量运算。',
         '本工作区只有 HAL_GetTick()（1kHz、1ms 分辨率），'
         '所有超时、周期、在线检测都基于它。',
         '1ms 分辨率做不了微秒级测量：帧间隔抖动、通信延迟、控制回路耗时统计都不行。'
         '另外 HAL_GetTick 依赖 SysTick 中断，抢占优先级最低，'
         '在长中断里会停止推进（HAL_Delay 在中断里用会出问题）。',
         '测量通信帧间隔、做精确超时（例如 IMU 超过 5ms 没数据就判定故障）、'
         '统计各任务实际耗时、给数据打时间戳做分析。',
         '🟢 完全可搬：F407 的 TIM2 与 TIM5 同样是 32 位定时器，做法一模一样，'
         '只需换 PSC/ARR 与时钟频率。')

    # -------- 3.9 工程规范与构建
    doc.add_heading('3.9 工程规范与构建系统', level=2)

    item(doc, 37, '目录分层 + 文件名前缀强制单向依赖',
         'User_File 分 1_Middleware（Algorithm / Driver / System）/ 2_Device（BSP / Motor / Plotter）/ '
         '5_Task，与 CubeMX 生成物物理隔离。文件名前缀严格自洽：'
         'alg_ 纯算法（零硬件依赖，alg_pid.h 只 include <stdint.h>，'
         '**可以原样拿到 PC 上做单元测试**）、sys_ 系统服务、drv_ 外设驱动、'
         'bsp_ 板级（绑引脚 / 器件 / 时序）、dvc_ 设备抽象（有协议有语义）、tsk_ 装配。'
         '调用链严格单向 alg_ ← dvc_/bsp_ ← drv_ ← tsk_，无任何跨层反向依赖。'
         '→ **看文件名就知道它能不能 include 谁**。',
         '本工作区分 Core / User_Drives / User_Architect / User_Algorithm / User_Application，'
         '分层本身清晰（本次移植就是照层次进行的），'
         '但文件名前缀不承载层次信息（user_can.c 在 User_Drives、user_pid.c 在 User_Algorithm，'
         '两者都叫 user_*），且 Algorithm 与 Archtect 的边界靠约定而非命名。',
         '参考工程的分层信息是"编码在文件名里"的，可以机械校验；'
         '本工作区是"编码在目录里"的，需要记住目录职责。'
         '另外参考工程的 alg_ 层零依赖这一点很有价值（可 PC 单测），'
         '本工作区的 user_coord.c 依赖 arm_math.h、user_pid.c 依赖 main.h，做不到。',
         '多人协作、跨赛季复用、需要给新人快速上手时。',
         '🟢 可搬（重构成本中等，但收益是长期的）。'
         '最小改动版：至少给控制器 ID 魔数换成枚举，并统一设备层命名。')

    item(doc, 38, 'Doxygen 骨架 + @note 记录"为什么这么做"',
         '每个文件统一的骨架：@file / @author / @brief / @version / @date + '
         'Includes / Exported macros / Exported types / Exported variables / '
         'Exported function prototypes / Exported function definitions。'
         '每个函数、每个成员变量都有 Doxygen 注释。'
         '**最有价值的是 @note 里写的"为什么"与"已知限制"**，例如：'
         '"防止 Pre_Encoder 默认 0 导致误判跨圈"、'
         '"积分项输出限幅的同时对积分误差也限幅，防止积分误差一直增大"、'
         '"这个功能也不能盲目开，比如 Pitch 轴电机为了克服重力…"、'
         '"当前版本仅完成人工上板功能测试并通过，尚未进行逐行代码审查、'
         '长期稳定性测试和异常注入测试"。'
         '→ 把踩过的坑固化在代码旁边，后人不会重复踩。',
         '本工作区是中文行注释 + 关键处块注释，'
         '注释质量本身不低（本次阅读中 user_uart.c 关于发送队列、'
         'user_coord.c 关于角度归一化的注释都很到位），'
         '但没有统一骨架，也没有系统性地记录"已知限制"。',
         '参考工程的注释体系把"设计决策"与"已知风险"变成了可检索的文档资产；'
         '本工作区更多是"就地解释这段代码在干什么"。',
         '知识传承、多人协作、自己几个月后回看。',
         '🟢 完全可搬（纯书写习惯）。建议至少对底层驱动与控制器补 @note 记录已知限制。')

    item(doc, 39, '构建系统：显式源文件列表 + 三种库类型 + CMSIS-DSP 聚合编译',
         '① **显式枚举源文件**（18 个 User_File 源文件 + 18 条 include 路径逐条列出），'
         '不用 file(GLOB)：新增文件必须显式登记，不会"写了文件却没编进去"，'
         'git diff 能看清源文件集合变更。'
         '② **三种库类型各得其所**：HAL 用 OBJECT 库（.o 直接并入可执行文件，'
         '规避 .a 归档与 --gc-sections 的粒度问题）、CMSIS-DSP 用 STATIC、'
         '接口宏用一个 INTERFACE 库通过 target_include_directories(INTERFACE) 与 '
         'target_compile_definitions(INTERFACE) 传播给使用方。'
         '③ **CMSIS-DSP 只编 7 个聚合文件**（BasicMathFunctions.c 这类文件内容就是 '
         '#include 掉同目录所有子 .c），并且按需裁剪（不编 Filtering / Transform / Complex）；'
         'ARM_MATH_CM7 用 PUBLIC 传播，保证使用方编译 inline 函数时走同一分支。'
         '④ C++ 编译选项 -fno-rtti -fno-exceptions -fno-threadsafe-statics 显著省体积；'
         '⑤ 链接加 -Wl,--print-memory-usage 每次打印 RAM/Flash 占用；'
         '⑥ 用 CMakePresets.json 管理 Debug/Release。',
         '本工作区是 file(GLOB_RECURSE) 一把梭（User_Drives / User_Algorithm / User_Architect / '
         'User_Application / User_Lib / SEGGER_RTT / Core / Drivers 八个通配），'
         '单个 add_executable，没有库目标划分，'
         '**CMSIS-DSP 是 564 个独立编译单元**（因为我们用的是 A 板带过来的那棵没有聚合文件的 DSP 树）。'
         '源码集合变化不用改 CMakeLists（这是 GLOB 的好处），但误引入文件也不会被发现。',
         '最实际的差距是**构建时间**：参考工程编 7 个 DSP 文件，本工作区编 564 个，'
         '首次全量构建差一个数量级。另外 GLOB 有一个已知坑：'
         '新增文件后必须重新 configure，否则新文件不参与编译（容易误判成"代码没生效"）。',
         '频繁全量构建、CI、多人反复 clone 时。',
         '🟢 完全可搬，且是**本报告里收益最直接的一条**。'
         '两种做法：① 在现有 DSP 树里新增 7 个聚合 .c（内容 #include 同目录子文件），'
         '并在 CMakeLists 里用 list(FILTER ... EXCLUDE REGEX) 把子文件排除；'
         '② 直接采用参考工程那棵自带聚合文件的 DSP 树。'
         '⚠️ 风险：若聚合后出现 static 函数重名冲突（不同子文件定义了同名 static 辅助函数），'
         '就退回做法 ② 或维持现状。')

    # ================= 四、本工作区的相对优势
    doc.add_heading('四、本工作区的相对优势（不要把好东西弄丢了）', level=1)
    para(doc, '对比必须是双向的。以下几条是本工作区明显更强、'
              '在借鉴参考工程时**不应该被动摇**的地方。', size=9.5, color=GREY)

    big_table(doc,
              ['优势', '本工作区怎么做', '参考工程现状'],
              [
                  ['控制器可任意组合（组合式架构）',
                   'CONTROLLER_INTERFACE 是一张函数指针虚表（Set_Target / Calculate / '
                   'Get_Output / Set_MaxOut / Get_MaxOut）。user_cascade.c 持有外环与内环两个 '
                   'CONTROLLER_INTERFACE* 指针；user_parallel.c 持有两个控制器指针 + 混合系数。'
                   '→ PID / 增量 PID / 容差 PID / ADRC / LADRC 可以任意嵌套、并联、串级组合，'
                   '而且电机是通过 DJI_Motor_Init(..., CONTROLLER_INTERFACE*) 注入控制器的。',
                   'Class_PID 是单一具体类；串级不在类内，而是在电机类里拼装'
                   '（PID_Angle 输出 → Target_Omega → PID_Omega → Target_Current）。'
                   '三个电机类各自内置这两个 PID，无法替换成别的控制器。'],
                  ['控制器种类多（7 种）',
                   'PID、增量 PID（带增量限幅）、容差 PID（带 tolerance）、'
                   'ADRC（跟踪微分器 + 非线性 ESO，参数含 r/h/beta/b0/dt）、'
                   'LADRC（线性 ESO，参数含观测器带宽 wc）、串级、并联。'
                   'ADRC / LADRC 都显式带 dt。',
                   '只有一种 PID（虽然特性很全）。要 ADRC 得自己写。'],
                  ['完整的 3D 数学库',
                   'user_coord.c/h：笛卡尔与极坐标两套坐标类型及全套运算（点积/叉积/距离/'
                   '夹角/归一化/投影/插值）、四元数（单位化、轴角构造、欧拉角互转、乘法、共轭、'
                   '模长、**Slerp**）、旋转矩阵（欧拉角/四元数互转、乘法、转置）、欧拉角、'
                   '坐标旋转与坐标系变换。这是姿态解算与云台控制的基础。',
                   'alg_basic 是低层数值小工具（端序、校验和、sinc、区间映射、'
                   '无效浮点判定、通用取模归一化、模板限幅）。'
                   '**没有向量、矩阵、四元数、插值、滤波、斜坡限制器。**'
                   '两者互补而非替代。'],
                  ['上层业务完整度',
                   '三种底盘模式（跟随 / 手动 / 小陀螺）+ 航向补偿 + 航向源可信度判断、'
                   '云台 yaw 锁世界 + 前馈、GM6020 绝对编码器航向源、开机音乐、'
                   'Serialplot 在线调参（带 min/max 钳位）、JScope 波形。'
                   '是一个可以上场的整车框架。',
                   '当前这份接入代码是"公共库 + 单电机验证"：1 个 C620、'
                   '正弦波形作角度目标、8 路 Serialplot 遥测。'
                   '底盘、云台、模式切换等业务尚未铺开。'],
                  ['串口/CAN 发送队列',
                   'user_uart.c 的 tx_queue 用节点空闲链表实现，可以连续排队多帧；'
                   '发送失败会把节点归还，不会堵死队列。'
                   'CAN 也有 tx_queue（Circular_Queue）。',
                   'UART 单槽 + HAL_BUSY、USB USBD_BUSY、CAN Tx FIFO 满则 HAL_BUSY，'
                   '**且上层调用一律忽略返回值** → 高负载静默丢帧、无计数无告警。'],
                  ['功率限制',
                   'DJI_Motor_Set_Power_Limit()，并在 chassis_config.h 定义 '
                   'CHASSIS_WHEEL_POWER_MAX 等参数，底盘 4 轮与云台都设了功率上限。',
                   'POWER_MODEL 四个系数全为 0 → power_calculate() 恒返回 0，'
                   'Get_Now_Power() 永远 0；唯一限幅是电调量程 ±16384。'
                   '**并且没有任何温度保护。**'],
                  ['纯 C + 静态内存',
                   '无 new/delete、无异常、无 RTTI，体积与最坏执行时间可预测；'
                   'F407 只有 1MB Flash / 192KB RAM，这个取舍是合理的'
                   '（当前实测占用 FLASH 4.2%、RAM 24%）。',
                   'C++17 类 / 模板 / 默认参数 / 引用 / 运算符重载。'
                   '代价是**算法层无法被 C 编译器解析**，要移植回 C 项目只能重写接口；'
                   '且三个电机类有约 600 行 × 3 的逐行重复。'],
                  ['可用按键 / 灯带 / 外部 Flash',
                   '（本工作区这几项没有；但反过来，参考工程也**没有**本工作区的 '
                   'BMI088 双片选 SPI、DBUS 9bit 校验遥控、HWT906 外部 IMU、'
                   'JScope RTT 波形。两边外设覆盖面不同，不是单方面领先。）',
                   '有按键（侵入式链表）、WS2812 灯带、W25Q64 外部 Flash、'
                   'USB CDC 虚拟串口、Vofa 协议库 —— 这些是本工作区没有的。'],
              ],
              [3.0, 7.6, 6.4])

    # ================= 五、差异汇总总表
    doc.add_heading('五、差异汇总总表（按可移植性分级）', level=1)

    doc.add_heading('5.1 🟢 可以直接搬（与 MCU 无关，纯算法 / 纯逻辑 / 纯规范）', level=2)
    big_table(doc,
              ['差异点', '参考工程', '本工作区', '工作量'],
              [
                  ['PID 显式 dt', '有 D_T，双重除零保护', '无，dt 隐含在 ki/kd 里', '小'],
                  ['微分先行', '有开关', '无', '小'],
                  ['D 项低通滤波', '一阶 IIR + alpha 钳位', '无', '小'],
                  ['D 项独立限幅', 'D_Out_Max', '无', '极小'],
                  ['变速积分', 'A/B 线性衰减', '无', '小'],
                  ['积分分离', '有（但硬清零，建议改成冻结）', '无', '小'],
                  ['死区', '有（算 P 之前清零误差）', '无（只在应用层做摇杆死区）', '极小'],
                  ['目标速度前馈 K_F', '有', '无（sub_feedback 形参预留未用）', '小'],
                  ['方向枚举', 'PID_DIRECT / PID_REVERSE', '无，调用者自己翻符号', '极小'],
                  ['零位积分泄放', '有（含适用反例注释）', '无', '小'],
                  ['全套 setter + Clear()', '15 个 Set_* + Clear()', '只有 Set_MaxOut', '小'],
                  ['无效浮点检测', 'Is_Invalid_Float（位级，含次正规数）', '无', '极小'],
                  ['端序 / 校验和 / 通用取模归一化', '有公共实现', '散落在各协议文件里', '小'],
                  ['接收时间戳 + 64 位微秒时间基准', 'TIM5 1MHz + 溢出计数 + 每帧打戳',
                   '只有 HAL_GetTick（1ms）', '中'],
                  ['main.c 极简化', '用户代码只有 Task_Init / Task_Loop',
                   'main.c 就是业务初始化本体', '小'],
                  ['CAN 发送放在定时器节拍里', '在 1ms 定时器中断里定周期发',
                   '在主循环 loop_event 里尽力而为', '小'],
                  ['Doxygen 骨架 + @note 记已知限制', '全工程统一', '中文行注释为主', '中（书写量）'],
                  ['CMSIS-DSP 聚合编译', '7 个编译单元', '564 个编译单元', '小（有风险）'],
                  ['显式源文件列表', '逐条枚举 + 三种库类型', 'file(GLOB_RECURSE) 一把梭', '中'],
                  ['文件名前缀强制层次', 'alg_/sys_/drv_/bsp_/dvc_/tsk_', '统一 user_*', '中（重构）'],
                  ['按键侵入式链表', '加按键 = 加一行', '无按键支持', '中（且需确认引脚）'],
                  ['蜂鸣器谱面结构', '{frequency, beat} 两字段 + 可读音名常量', '单曲开机音乐', '中'],
              ],
              [4.0, 5.3, 5.3, 1.8])

    doc.add_heading('5.2 🟡 需要改造（依赖体系结构或外设能力）', level=2)
    big_table(doc,
              ['差异点', '说明'],
              [
                  ['电机按电调建类 + 偏移指针共帧',
                   '本工作区已有可用的统一组帧实现，收益主要是省 CPU 与可读性；'
                   '建议只做其中的低成本部分：**电机 ID 枚举化**、'
                   '**确认首帧编码器基准保护**、**确认共帧残留电流清零**。'],
                  ['Manage_Object 全局实例表',
                   '本工作区的运行时注册表也能工作，且加实例只改一个宏；'
                   '不必为了"像参考工程"而推翻，但如果要加很多路 UART 就会撞上限（UART_NUM=5）。'],
                  ['外部反馈覆盖 + 单次有效',
                   '本工作区的 CONTROLLER_INTERFACE 注入机制能力等价甚至更直接，不需要移植。'],
                  ['Vofa 协议（justfloat + 累加和）',
                   '纯协议逻辑，改动小；但价值取决于是否使用 VOFA+ 上位机。'
                   '注意参考工程这个模块"完备但零接入"。'],
                  ['绘图协议可选数据类型 / 可关校验',
                   '纯协议逻辑，改动小。本工作区的固定 float32 + 钳位调参表其实够用且更安全。'],
                  ['波形发生器升级为辨识激励台',
                   '波形算法两边都有；缺的是"给激励 → 采响应 → 拟合参数"的工作流，'
                   '属于补文档/脚本而非补代码。'],
                  ['按键 / WS2812 / W25Q64 外部 Flash',
                   '都是**参考工程多出来的外设能力**，属于新功能而非移植。'
                   'C 板上要先确认有没有对应可用引脚与外设。'],
                  ['中间件抽成可复用公共库',
                   '参考工程把 1_Middleware 明确按"长期复用、稳定优先、兼容优先"维护。'
                   '本工作区的分层在同工程内是清晰的，但未抽出跨赛季复用的库边界。'],
              ],
              [5.0, 11.8])

    doc.add_heading('5.3 🔴 搬不过来（H7 专属，F407 上不存在）', level=2)
    big_table(doc,
              ['H7 专属特性', 'F407 对应物 / 处理方式'],
              [
                  ['FDCAN（drv_can 全部）',
                   'F407 是 bxCAN：过滤器组 / bank、FIFO0/1、邮箱、CAN_FilterTypeDef，'
                   '接口形态完全不同。本工作区的 user_can.c 已经是 bxCAN 版本，无需改动。'],
                  ['BDMA（SPI6 用）',
                   'F407 无 BDMA；对应物是 DMA1/DMA2。C 板 SPI1 用普通 DMA 即可。'],
                  ['MDMA（OSPI 搬运）',
                   'F407 无 MDMA，只有 DMA1/DMA2。'],
                  ['OCTOSPI + OCTOSPIM + DLYB/DHQC',
                   'F407 只有 QUADSPI（无 OSPI/OSPIM/DLYB/DHQC/HAL_OSPI_*）。'],
                  ['Cortex-M7 Cache 维护',
                   'F407 是 M4，无 Cache、无 SCB_CCR_DC、无 SCB_CleanDCache_by_Addr。'
                   '（注：参考工程本身也只开了 I-Cache、D-Cache 关闭，'
                   '所以那些 cache 维护函数目前其实都在空转。）'],
                  ['多域 SRAM（DTCM / RAM_D2 / RAM_D3）',
                   'F407 的对应物是 CCM RAM（同样 DMA 不可访问）；F407 的普通 SRAM 全部可被 DMA 访问，'
                   '不存在"按外设所在电源域选 DMA 控制器"的复杂度。'],
                  ['每外设独立内核时钟树 + VOS0 + PLLFRACN 小数分频',
                   'F407 无 HAL_RCCEx_PeriphCLKConfig 的 D1/D2/D3 分频、无 VOS、'
                   'PLL 无小数分频（FRACN）。本工作区的时钟树简单得多，也不该去模仿。'],
                  ['UART 数量与 FIFO API',
                   'F407 只有 USART1/2/3/6 + UART4/5，且 USART 无 FIFO，'
                   'HAL_UARTEx_SetTxFifoThreshold 等 API 不存在。'],
                  ['USB_OTG_HS + 内嵌 FS PHY + HSI48',
                   'F407 的 USB OTG FS 时钟源是 PLL48CK，PCD 初始化与中断向量都要改。'],
                  ['「H7 双核」',
                   '**不存在**：H723 是单核 Cortex-M7，工程里没有任何 HSEM / IPCC / OpenAMP 调用。'
                   'stm32h7xx_hal_conf.h 里的 HAL_HSEM_MODULE_ENABLED 只是 CubeMX 模板残留。'],
              ],
              [5.6, 11.2])

    # ================= 六、落地建议
    doc.add_heading('六、落地建议（按性价比排序）', level=1)
    para(doc, '每一项都标注了大致工作量与风险，可以直接当 TODO 用。', size=9.5, color=GREY)

    big_table(doc,
              ['优先级', '事项', '做法', '收益', '工作量 / 风险'],
              [
                  ['★★★', 'PID 工程化特性移植',
                   '给 PID_Controller 加 dt、Iout 乘 dt、Dout 除 dt；加 D 项限幅、'
                   'D 项一阶低通、微分先行、死区、变速积分、前馈、方向枚举、'
                   'PID_Clear()；补齐全套 setter。保持 C 结构体风格，不要引入 C++。',
                   '唯一直接影响控制效果的部分；抗超调、抗噪声、抗阶跃冲击、'
                   '支持在线整定。全部与 MCU 无关。',
                   '半天。风险低。'
                   '注意：积分分离要改成"冻结/衰减"而不是参考工程的硬清零。'],
                  ['★★★', 'CMSIS-DSP 聚合编译',
                   '在 Drivers/CMSIS/DSP/Source 各子目录新增 7 个聚合 .c'
                   '（内容为 #include 同目录子文件），并在 CMakeLists 里用 '
                   'list(FILTER ... EXCLUDE REGEX) 排除子文件。',
                   '编译单元 564 → 7，首次全量构建时间数量级下降；'
                   '每次改配置的重配置也更快。',
                   '1 小时。风险中：若出现 static 重名冲突则改用参考工程那棵'
                   '自带聚合文件的 DSP 树，或维持现状。'],
                  ['★★☆', '新增无效浮点数检测',
                   '写一个 Is_Invalid_Float（位级：指数 0 且尾数非 0 即次正规数判无效；'
                   '指数全 1 即 Inf/NaN 判无效），在 BMI088 采样、DBUS 解析、'
                   '串口调参入口各加一道，异常则丢弃该帧并计数。',
                   '防御"脏值一路流到电机电流指令"这种最危险也最难查的故障。',
                   '1 小时。风险极低。'],
                  ['★★☆', 'CAN 发送挪进定时器节拍',
                   '把 CAN_QueueHandle 从 loop_event 主循环表里挪到一个定时器回调'
                   '（例如复用 JScope 的 TIM10，或新增一个 1ms 定时器任务）。',
                   '发送时序从"尽力而为"变成"确定周期"，多电机同步下发更稳。',
                   '1 小时。风险低。注意控制 CAN 发送频率与总线负载。'],
                  ['★★☆', '接收时间戳 + 64 位微秒时间基准',
                   '用 F407 的 TIM2 或 TIM5（32 位）跑 1MHz + 溢出计数，'
                   '提供 Get_Timestamp_us()；给 UART_DRIVES / CAN_DRIVES 加接收时间戳字段。',
                   '能判断"多久没来数据"、测帧间隔抖动、做精确超时。',
                   '半天。风险低。'],
                  ['★★☆', 'main.c 极简化',
                   '把 main.c 的 USER CODE 2 内容搬进 User_Application/app_init.c 的 '
                   'App_Init() / App_Loop()，main.c 只留两行调用。',
                   '对 CubeMX 重新生成结构性免疫；main.c 变成一块可读的"接线板"。',
                   '半天。风险低，但改动面较大，建议在功能稳定期做。'],
                  ['★☆☆', 'PID 增益在线可调 + 调参表安全化',
                   '补齐全套 setter；把 Serialplot 调参表与 PID 参数绑定，'
                   '并保留现有的 min/max 钳位。',
                   '可以在跑车过程中整定，不用反复烧写。',
                   '2 小时。风险低。'],
                  ['★☆☆', '注释体系规范化',
                   '至少对 User_Drives 与 User_Algorithm 补 Doxygen 骨架；'
                   '重点是用 @note 记录"已知限制"与"为什么这么设计"。',
                   '把踩过的坑变成可检索的文档资产。',
                   '1~2 天（书写量）。风险无。'],
                  ['★☆☆', '电机 ID 枚举化 + 确认两处保护',
                   '把 DJI_Motor_Init 里的裸数字 2/3/1/4 换成 Enum_Motor_DJI_ID 风格的枚举；'
                   '核对首帧编码器基准保护、共帧残留电流清零是否与参考工程等价。',
                   '可读性 + 两个潜在偶发 bug 的排除。',
                   '2 小时。风险低。'
                   '**其中"确认两处保护"优先级应当提高**：若本工作区缺失，属于潜在 bug。'],
              ],
              [1.4, 3.0, 4.9, 3.9, 3.9])

    # ================= 七、参考工程自身要注意的坑
    doc.add_heading('七、参考工程自身需要注意的坑（避免盲目照抄）', level=1)
    para(doc, '这一章是本次阅读中发现的、参考工程自己的问题。'
              '借鉴它的设计时要知道这些，否则会把 bug 一起搬过来。', size=9.5, color=GREY)

    big_table(doc,
              ['问题', '证据', '后果 / 借鉴时的注意'],
              [
                  ['PID 特性"齐备但未启用"',
                   'tsk_config_and_callback.cpp 里 PID_Omega.Init(0,0,0) / PID_Angle.Init(0,0,0) '
                   '只传了 3 个参数，其余 13 个全走默认值。',
                   '三个限幅全为 0（不限制）、无死区、无滤波、无变积分、ZPIB 关闭 —— '
                   '第 3.1 节描述的特性在该构建里**一个都没生效**，唯一限幅是设备层的 ±16384。'
                   '借鉴时要明确：**功能存在 ≠ 功能已启用**，移植后必须显式配置每个参数。'],
                  ['Clear() 之后首拍微分/前馈冲击',
                   'Clear() 把 Pre_Now / Pre_Target 置 0，没有"首拍初始化"机制。',
                   '开微分先行或前馈时，使能首拍会产生 -K_D·Now/dt 量级的满量程冲击'
                   '（按其注释算例约 −500）。'
                   '**移植时必须加首拍保护**：第一拍令 Pre_Now = Now、Pre_Target = Target。'],
                  ['积分分离是硬清零而非冻结',
                   '源码里超阈值时直接 Integral_Error = 0。',
                   '对"靠积分顶重力"的轴（云台 pitch）大阶跃瞬间会掉头/下沉 —— '
                   '与其自身 ZPIB 注释里"Pitch 轴需积分顶重力"的说法自相矛盾。'
                   '移植时应改成冻结或缓慢衰减。'],
                  ['变积分与积分分离会互相吃掉',
                   '若 I_Separate_Threshold < A + B，误差超阈值时积分被清零，'
                   '永远进不了 [A, A+B] 过渡区。',
                   '变积分变成死代码，且两参数各自合法、组合失效、无断言无文档。'
                   '移植时要么二选一，要么加参数合法性校验。'],
                  ['ZPIB 与 Dead_Zone 隐式耦合',
                   'ZPIB 触发条件含 |Target| < Dead_Zone。',
                   'Dead_Zone = 0 时 ZPIB 永不触发（开了也没用）。'
                   '移植时要显式说明这个依赖。'],
                  ['无角度环绕',
                   'Error = Target - Now 不做归一化，隔壁的 Modulus_Normalization 没被调用。',
                   '目前靠多圈累加侥幸不暴露；一旦换成 ±π 归一化角度（云台 yaw 常见）'
                   '就会跨 ±π 绕远路。本工作区在做角度环时**必须自己做环绕**。'],
                  ['无 NaN / Inf 防护',
                   'Is_Invalid_Float 已实现，但全工程零调用。',
                   '上游坏帧产生 NaN 后积分累加器被永久污染，只有人工 Clear() 能恢复。'
                   '这正好说明了本报告 3.2 第 12 条的价值 —— 工具写了不用等于没写。'],
                  ['drv_spi 回调长度恒为 0（真 bug）',
                   'HAL_SPI_TxCpltCallback / TxRxCpltCallback 先调用清状态函数'
                   '（把长度清零）再回调，所以回调收到的长度永远是 0。'
                   '对比 drv_ospi 是"先存局部变量再清状态"，写法正确。',
                   '目前唯一使用者（WS2812）传的是空回调，所以没暴露。'
                   '**若本工作区将来要加 SPI 的 DMA 回调，注意顺序**。'],
                  ['中断优先级倒挂',
                   '所有通信外设抢占优先级都是 0，而 1ms 调度器（TIM7）是 5。',
                   'CAN 接收风暴会持续抢占 1ms 任务，导致控制周期抖动。'
                   '本工作区当前是 CAN/UART/TIM10 = 0、TIM4 = 2、SysTick = 15，'
                   '同样存在"调度器被通信抢占"的问题，规划优先级时要注意。'],
                  ['Delay_* 忙等 + 时基靠中断 → 可死锁',
                   '时间戳由 TIM5 中断（优先级 5）推进，而 Delay_* 是忙等读时间戳。'
                   '在任何抢占优先级数值小于 5 的中断里调用它就会死锁。',
                   '当前调用点在主循环所以没踩到，但没有约束也没有保护性注释。'
                   '若本工作区引入微秒延时，**必须禁止在中断里调用**。'],
                  ['init_finished 是普通 bool 却在 ISR 里读',
                   'drv_can / drv_uart / drv_usb / drv_tim 的回调里都在读它。',
                   '严格意义的数据竞态。移植时应写成 volatile。'],
                  ['USB 绘图发送数据竞争',
                   'CDC_Transmit_HS 只记指针不拷贝，而 1ms 任务每周期覆写同一块缓冲，'
                   '且返回值被忽略。UART 版有 gState 保护、USB 版没有。',
                   '会间歇发出撕裂帧。本工作区走 UART + 队列，没有这个问题，'
                   '**不要**为了"支持 USB 绘图"而照抄这个实现。'],
                  ['蜂鸣器 Play_Note 用 HAL_Delay 阻塞',
                   '在主循环里播一首曲子会阻塞数十秒，靠"只播一次"掩盖。',
                   '任何"运行中报警响一声"都会卡住主循环。'
                   '正确做法是 1ms tick 驱动的 {freq, duration} 状态机。'],
                  ['W25Q64 自动轮询无软件超时',
                   '头文件注释自认"自动轮询当前未配置软件超时，异常情况下可能长期保持 Busy"。',
                   'Flash 若一直 BUSY，后续所有操作永久 HAL_BUSY，且应用层无法区分'
                   '"暂时忙"与"永久坏"。工程已有 64 位微秒时间戳，加超时成本极低。'],
                  ['默认 Debug(-O0) 构建',
                   'CMakePresets 默认 Debug，CXX_FLAGS_DEBUG = "-O0 -g3"，'
                   '而电机头文件里有 1803 行 inline getter/setter（-O0 下不内联）。',
                   '用 -O0 跑 1kHz 中断是危险组合；README 也没提示"上机必须用 Release"。'
                   '本工作区默认 -Og -g，比 -O0 好一些，但同样建议上机用 Release/MinSizeRel。'],
                  ['.vscode/tasks.json 硬编码他机路径',
                   '写死了原开发机的 CLion cmake.exe、C:/0Embedded/... 的 build 目录与 OpenOCD，'
                   '没有一处用 ${workspaceFolder}。',
                   '新队员 clone 后任务直接失败。本工作区刚刚修好了对应的任务'
                   '（改用相对路径 + 缺配置时自动配置），**不要**把它改回硬编码风格。'],
                  ['WS2812 头文件注释与代码不一致',
                   '注释写 0xC0 / 0xF0，代码实际是 0x60 / 0x78（整体右移一位），'
                   '且注释里"低电平 1000ns"应为 833ns。',
                   '代码是对的、注释是错的。若按注释"修正"会引入半时钟相位错误。'],
                  ['三个电机类约 600 行 × 3 逐行重复',
                   'GM6020 / C610 / C620 互不继承，Init / CAN_RxCpltCallback / '
                   'TIM_Calculate / PID_Calculate / Output / Clear 几乎完全相同。',
                   '改一个 bug 要改三处，极易漏改。'
                   '本工作区只有一个统一的 DJI_MOTOR_DRIVES，**这一点本工作区结构更好**，'
                   '不要为了"像参考工程"而拆成三份。'],
              ],
              [3.6, 5.6, 7.6])

    # ================= 八、附录
    doc.add_heading('八、附录：核对的文件与关键位置', level=1)
    big_table(doc,
              ['层次', '参考工程文件', '本工作区对应文件'],
              [
                  ['控制器算法', 'User_File/1_Middleware/Algorithm/PID/alg_pid.h/.cpp',
                   'User_Algorithm/User_Controller/user_pid.h + Src/user_pid.c'],
                  ['控制器家族（本工作区独有）', '（无）',
                   'user_inc_pid / user_pid_tolerance / user_adrc / user_ladrc / '
                   'user_cascade / user_parallel / user_controller.h'],
                  ['基础数学', 'Algorithm/Basic/alg_basic.h/.cpp、alg_arm_math.h',
                   'User_Algorithm/user_coord.h/.c、user_cf.c、user_fir.c、user_newton.c'],
                  ['波形发生器', 'Algorithm/Waveform/alg_waveform.h/.cpp',
                   'User_Algorithm/user_waveform.h/.c'],
                  ['外设驱动', 'Middleware/Driver/{CAN,UART,SPI,TIM,USB,OSPI}/drv_*.h/.cpp',
                   'User_Drives/User_Peripheral/user_{can,uart,spi,pwm,timer}.h/.c'],
                  ['时间基准', 'Middleware/System/Timestamp/sys_timestamp.h/.cpp',
                   '（无对应模块，仅 HAL_GetTick）'],
                  ['DJI 电机', 'Device/Motor/Motor_DJI/dvc_motor_dji.h/.cpp',
                   'User_Drives/User_Motor/user_dji_motor.h + Src/user_dji_motor.c'],
                  ['上位机绘图', 'Device/Plotter/Serialplot/dvc_serialplot.*、Vofa/dvc_vofa.*',
                   'User_Drives/User_Peripheral/user_serialplot.h/.c'],
                  ['遥控 / IMU', '（无对应模块）',
                   'User_Drives/user_dji_bus.h/.c、User_Sensor/user_hwt906.*、user_bmi088.*'],
                  ['BSP', 'Device/BSP/{Buzzer,Key,WS2812,W25Q64JV}/bsp_*',
                   'User_Drives/User_Basic/user_{led,buzzer}.h/.c'],
                  ['姿态解算', '（无对应模块）', 'User_Algorithm/User_Attitude/attitude.h + Src/attitude.c'],
                  ['任务装配', 'User_File/5_Task/tsk_config_and_callback.h/.cpp',
                   'Core/Src/main.c 的 USER CODE 2 + User_Architect/user_systick.*'],
                  ['分层/构建', 'CMakeLists.txt、CMakePresets.json、.vscode/tasks.json、Sentry.ioc',
                   'CMakeLists.txt、CMakePresets.json、.vscode/tasks.json、Desktop.ioc'],
              ],
              [3.0, 7.4, 6.4])

    doc.add_paragraph()
    para(doc, '报告结束。本文档基于对两个工程源码的逐文件阅读生成，'
              '所有对比结论均可在上述文件中按函数名 / 行号复核。',
         size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(path)
    return path


if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else 'compare_report.docx'
    build(out)
    print('generated:', out)
